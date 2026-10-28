#!/usr/bin/env python3
"""Create an editable report scaffold. This does not write a finished report."""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


def font(style, size=14, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), "Times New Roman")
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)
    complex_size = props.find(qn("w:szCs"))
    if complex_size is not None:
        complex_size.set(qn("w:val"), str(size * 2))


def field(paragraph, instruction, cached=""):
    run = paragraph.add_run()
    for kind in ("begin", "separate", "end"):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
        if kind == "begin":
            text = OxmlElement("w:instrText")
            text.set(qn("xml:space"), "preserve")
            text.text = " " + instruction + " "
            run._r.append(text)
        elif kind == "separate":
            text = OxmlElement("w:t")
            text.text = cached
            run._r.append(text)


def configure(doc, has_title):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(30), Mm(15)
    sec.top_margin = sec.bottom_margin = Mm(20)
    sec.footer_distance = Mm(10)
    sec.different_first_page_header_footer = has_title
    base = doc.styles["Normal"]
    font(base)
    pf = base.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.widow_control = True
    for level in range(1, 4):
        st = doc.styles[f"Heading {level}"]
        st.base_style = base
        font(st, bold=True)
        p = st.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.first_line_indent = Cm(1.25)
        p.space_before, p.space_after = Pt(12 if level > 1 else 0), Pt(12)
        p.line_spacing = 1.5
        p.keep_with_next = p.keep_together = True
        p.page_break_before = level == 1
    for name in ("Report Structure", "Report Front"):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = base
        font(st, bold=True)
        p = st.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.first_line_indent = Cm(0)
        p.space_after = Pt(12)
        p.page_break_before = True
        p.keep_with_next = True
        if name == "Report Structure":
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            st.element.get_or_add_pPr().append(outline)
    for name in ("Title", "TOC Heading", "Caption", "Footer"):
        st = doc.styles[name]
        st.base_style = base
        font(st, size=14, bold=name in ("Title", "TOC Heading"))
        for residue in st.element.xpath("./w:pPr/w:pBdr | ./w:rPr/w:u"):
            residue.getparent().remove(residue)
        p = st.paragraph_format
        p.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1 if name == "Caption" else 1.5
        p.keep_with_next = name != "Caption"
    footer = sec.footer.paragraphs[0]
    footer.style = doc.styles["Footer"]
    # The supplied document is a fragment following a separate title page.
    # The contents page is the first physical page and visibly numbered 2.
    first_number = 1 if has_title else 2
    field(footer, "PAGE", str(first_number))
    start = OxmlElement("w:pgNumType")
    start.set(qn("w:start"), str(first_number))
    sec._sectPr.append(start)
    update = doc.settings.element.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        doc.settings.element.append(update)
    update.set(qn("w:val"), "true")


def title_page(doc, meta, kind):
    labels = {"lr": "по лабораторной работе", "dz": "по домашнему заданию",
              "review": "по учебному заданию", "nirs": "по научно-исследовательской работе студента"}
    for key, default in (("university", "[Полное наименование вуза]"),
                         ("faculty", "[Факультет]"),
                         ("department", "[Кафедра]"),
                         ("program", "[Направление подготовки]")):
        p = doc.add_paragraph(str(meta.get(key, default)))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    p = doc.add_paragraph("ОТЧЕТ", "Title")
    p.paragraph_format.space_before = Pt(55)
    suffix = " № " + str(meta["number"]) if meta.get("number") else ""
    p = doc.add_paragraph(labels[kind] + suffix)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    if meta.get("variant"):
        p = doc.add_paragraph("Вариант № " + str(meta["variant"]))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    p = doc.add_paragraph(str(meta.get("title", "[Тема работы]")), "Title")
    p.paragraph_format.space_before = Pt(18)
    p = doc.add_paragraph("Дисциплина: " + str(meta.get("course", "[Дисциплина]")))
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(6), Cm(10.5)
    rows = [("Студент", str(meta.get("student", "[ФИО]"))),
            ("Группа", str(meta.get("group", "[Группа]"))),
            ("Преподаватель", str(meta.get("teacher", "[ФИО преподавателя]")))]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(6)
    p = doc.add_paragraph(str(meta.get("city", "[Город]")) + " " +
                          str(meta.get("year", "[Год]")))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(24)


def create_report(output, meta, kind="lr", toc=True, no_title=True):
    doc = Document()
    configure(doc, not no_title)
    if not no_title:
        title_page(doc, meta, kind)
    if toc:
        p = doc.add_paragraph("СОДЕРЖАНИЕ", "TOC Heading")
        p.paragraph_format.page_break_before = bool(doc.paragraphs[:-1])
        field(doc.add_paragraph(), 'TOC \\o "1-3" \\t "Report Structure,1" \\h \\z',
              "[Обновить оглавление после заполнения и проверить страницы]")
    p = doc.add_paragraph("ВВЕДЕНИЕ", "Report Structure")
    if no_title and not toc and kind != "nirs":
        p.paragraph_format.page_break_before = False
    doc.add_paragraph("[Цель, задачи, исходные условия и ограничения]")
    plans = {
        "lr": ["Теоретическая часть", "Практическая часть", "Обработка и анализ результатов"],
        "dz": ["Исходные данные и расчет", "Исследование модели", "Анализ результатов"],
        "review": ["Теоретические основы", "Реализация и применение", "Сравнение подходов и ограничения"],
        "nirs": ["Теоретические основы", "Сравнение и выбор подхода", "Методика и результаты", "Ограничения"]}
    for num, heading in enumerate(plans[kind], 1):
        doc.add_paragraph(f"{num} {heading}", "Heading 1")
        doc.add_paragraph("[Заполнить по заданию и фактическим материалам]")
    doc.add_paragraph("ЗАКЛЮЧЕНИЕ", "Report Structure")
    doc.add_paragraph("[Заполнить и проверить]")
    # Add a bibliography only when needed, using verified public sources.
    # Add a NIRS abstract only when required by the current assignment.
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--kind", choices=("lr", "dz", "review", "nirs"), default="lr")
    parser.add_argument("--metadata", type=Path, help="UTF-8 JSON object; omitted fields become markers")
    parser.add_argument("--output", type=Path, required=True)
    contents = parser.add_mutually_exclusive_group()
    contents.add_argument("--toc", dest="toc", action="store_true")
    contents.add_argument("--no-toc", dest="toc", action="store_false",
                          help="omit contents only when explicitly requested")
    title = parser.add_mutually_exclusive_group()
    title.add_argument("--title", dest="no_title", action="store_false",
                       help="include a title page only when explicitly required")
    title.add_argument("--no-title", dest="no_title", action="store_true")
    parser.set_defaults(toc=True, no_title=True)
    args = parser.parse_args()
    meta = json.loads(args.metadata.read_text(encoding="utf-8")) if args.metadata else {}
    if not isinstance(meta, dict):
        parser.error("metadata must contain a JSON object")
    print(create_report(args.output, meta, args.kind, args.toc, args.no_title))


if __name__ == "__main__":
    main()
